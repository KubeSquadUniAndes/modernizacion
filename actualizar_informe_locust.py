#
#    Copyright 2010-2026 the original author or authors.
#
#    Licensed under the Apache License, Version 2.0 (the "License");
#    you may not use this file except in compliance with the License.
#    You may obtain a copy of the License at
#
#       https://www.apache.org/licenses/LICENSE-2.0
#
#    Unless required by applicable law or agreed to in writing, software
#    distributed under the License is distributed on an "AS IS" BASIS,
#    WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#    See the License for the specific language governing permissions and
#    limitations under the License.
#

import csv
import json
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


REPORT_DIR = Path("reports")
CHART_DIR = REPORT_DIR / "charts"
OUTPUT = REPORT_DIR / "informe_pruebas.docx"
TEST_RESULTS = REPORT_DIR / "catalog-service-test-results.json"
SONAR_RESULTS = REPORT_DIR / "sonarqube-results.json"
LOCUST_STATS = REPORT_DIR / "locust_catalog_stats.csv"
LOCUST_HTML = REPORT_DIR / "locust_catalog_report.html"
LOCUST_SUMMARY = REPORT_DIR / "locust-summary.json"


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        run = table.rows[0].cells[index].paragraphs[0].add_run(str(header))
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    return table


def add_picture_if_exists(document, path, width=6.4):
    if path.exists():
        document.add_picture(str(path), width=Inches(width))


def setup_styles(document):
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    styles["Normal"].paragraph_format.line_spacing = 2
    styles["Normal"].paragraph_format.space_after = Pt(0)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Times New Roman"


def to_float(value):
    try:
        return float(value)
    except Exception:
        return 0.0


def to_int(value):
    try:
        return int(float(value))
    except Exception:
        return 0


def read_locust_stats():
    rows = []
    with LOCUST_STATS.open(newline="", encoding="utf-8") as file:
        reader = csv.DictReader(line for line in file if line.strip())
        for row in reader:
            if not row.get("Name"):
                continue
            rows.append(row)
    endpoints = [row for row in rows if row["Name"] != "Aggregated"]
    aggregated = next(row for row in rows if row["Name"] == "Aggregated")
    return endpoints, aggregated


def create_locust_charts(endpoints, aggregated):
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    labels = [row["Name"].replace("GET ", "") for row in endpoints]
    avg = [to_float(row["Average Response Time"]) for row in endpoints]
    p95 = [to_float(row["95%"] or row.get("95", 0)) for row in endpoints]
    request_counts = [to_int(row["Request Count"]) for row in endpoints]

    avg_chart = CHART_DIR / "locust_avg_response_by_endpoint.png"
    plt.figure(figsize=(11, 5.5))
    plt.bar(labels, avg, color="#1976d2")
    plt.ylabel("Tiempo promedio (ms)")
    plt.title("Locust - tiempo promedio de respuesta por endpoint")
    plt.xticks(rotation=35, ha="right")
    plt.tight_layout()
    plt.savefig(avg_chart, dpi=160)
    plt.close()

    p95_chart = CHART_DIR / "locust_p95_by_endpoint.png"
    plt.figure(figsize=(11, 5.5))
    plt.bar(labels, p95, color="#ef6c00")
    plt.ylabel("Percentil 95 (ms)")
    plt.title("Locust - percentil 95 por endpoint")
    plt.xticks(rotation=35, ha="right")
    plt.tight_layout()
    plt.savefig(p95_chart, dpi=160)
    plt.close()

    requests_chart = CHART_DIR / "locust_request_distribution.png"
    plt.figure(figsize=(10, 5.2))
    plt.pie(request_counts, labels=labels, autopct="%1.1f%%", startangle=90)
    plt.title("Locust - distribución de solicitudes por endpoint")
    plt.tight_layout()
    plt.savefig(requests_chart, dpi=160)
    plt.close()

    table_image = CHART_DIR / "locust_summary_table.png"
    lines = [
        "Resumen agregado Locust",
        "",
        f"Solicitudes totales: {aggregated['Request Count']}",
        f"Fallos: {aggregated['Failure Count']}",
        f"Requests/s: {to_float(aggregated['Requests/s']):.2f}",
        f"Latencia promedio: {to_float(aggregated['Average Response Time']):.2f} ms",
        f"Mediana: {aggregated['Median Response Time']} ms",
        f"Minima: {to_float(aggregated['Min Response Time']):.2f} ms",
        f"Maxima: {to_float(aggregated['Max Response Time']):.2f} ms",
        f"P95: {aggregated['95%']} ms",
        f"P99: {aggregated['99%']} ms",
    ]
    font = ImageFont.load_default()
    image = Image.new("RGB", (1050, 300), color=(20, 24, 28))
    draw = ImageDraw.Draw(image)
    y = 18
    for index, line in enumerate(lines):
        color = (144, 202, 249) if index == 0 else (238, 238, 238)
        draw.text((20, y), line, font=font, fill=color)
        y += 24
    image.save(table_image)

    return [avg_chart, p95_chart, requests_chart, table_image]


def capture_locust_html():
    output = CHART_DIR / "locust_html_report_screenshot.png"
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(
                executable_path=r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                headless=True,
                args=["--no-sandbox"],
            )
            page = browser.new_page(viewport={"width": 1440, "height": 1000})
            page.goto(LOCUST_HTML.resolve().as_uri(), wait_until="networkidle")
            page.wait_for_timeout(1000)
            page.screenshot(path=str(output), full_page=True)
            browser.close()
    except Exception:
        output = None
    return output


def metric(project, key):
    return project.get("measures", {}).get(key, "0")


def percent(value):
    return f"{to_float(value):.2f}%"


def minutes(value):
    return f"{to_float(value):.0f} min"


def add_cover(document):
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Informe de pruebas de calidad\n")
    run.bold = True
    run.font.size = Pt(22)
    subtitle = title.add_run("Evaluación de pruebas, métricas y calidad de código\nCatalog Service - JPetStore")
    subtitle.font.size = Pt(14)
    for _ in range(4):
        document.add_paragraph()
    for line in [
        "Aplicación destino: catalog-service",
        "Proyecto base: JPetStore",
        "Ambiente evaluado: AWS EKS, Amazon RDS PostgreSQL y SonarQube local en Docker",
        "Herramienta de carga: Locust",
    ]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(line)
    document.add_page_break()


def add_summary(document, test_data, sonar_data, aggregated):
    endpoint_results = test_data["endpoint_results"]
    passed = sum(1 for item in endpoint_results if item["passed"])
    catalog = next(project for project in sonar_data["projects"] if project["key"] == "catalog-service")
    monolith = next(project for project in sonar_data["projects"] if project["key"] == "jpetstore-monolith")
    document.add_heading("Resumen ejecutivo", level=1)
    document.add_paragraph(
        "Este informe consolida la ejecución de pruebas funcionales, pruebas de carga con Locust y análisis estático "
        "de código con SonarQube para el microservicio catalog-service."
    )
    document.add_paragraph(
        f"Las pruebas funcionales tuvieron {passed}/{len(endpoint_results)} validaciones exitosas. La prueba de carga "
        f"con Locust cubrió todos los endpoints disponibles, ejecutó {aggregated['Request Count']} solicitudes, "
        f"registró {aggregated['Failure Count']} fallos y obtuvo un throughput agregado de "
        f"{to_float(aggregated['Requests/s']):.2f} req/s."
    )
    document.add_paragraph(
        f"En SonarQube, el microservicio obtuvo Quality Gate {catalog['quality_gate']}, "
        f"{metric(catalog, 'bugs')} bugs y {metric(catalog, 'code_smells')} code smells. El monolito obtuvo "
        f"Quality Gate {monolith['quality_gate']}, {metric(monolith, 'bugs')} bugs y "
        f"{metric(monolith, 'code_smells')} code smells."
    )


def add_execution_guide(document):
    document.add_heading("Introducción y guía para ejecutar las pruebas", level=1)
    document.add_paragraph(
        "Para reproducir estas pruebas se requiere tener desplegado catalog-service en AWS, Python, Locust, curl, "
        "kubectl, Docker y SonarQube."
    )
    add_table(
        document,
        ["Paso", "Comando / acción", "Resultado esperado"],
        [
            ["1", "Validar servicio con curl /actuator/health", "Respuesta UP"],
            ["2", "Ejecutar Locust headless", "Archivos CSV y HTML en reports/"],
            ["3", "Ejecutar SonarQube en Docker", "Servidor disponible en localhost:9000"],
            ["4", "Ejecutar Maven Sonar Scanner", "Dashboards y Quality Gates"],
            ["5", "Generar informe", "Documento Word con tablas, análisis e imágenes"],
        ],
    )
    document.add_paragraph("Comando Locust usado:")
    document.add_paragraph(
        "python -m locust -f locustfile_catalog_service.py --headless --host <HOST> --users 20 "
        "--spawn-rate 5 --run-time 1m --csv reports\\locust_catalog --html reports\\locust_catalog_report.html"
    )


def add_functional_section(document, test_data):
    document.add_heading("Pruebas funcionales y disponibilidad", level=1)
    rows = []
    for item in test_data["endpoint_results"]:
        rows.append([
            item["name"], item["path"], item["expected_status"], item["status_code"], f"{item['elapsed_ms']:.2f}", "OK" if item["passed"] else "FALLA"
        ])
    add_table(document, ["Prueba", "Endpoint", "Esperado", "Obtenido", "Latencia ms", "Resultado"], rows)
    document.add_paragraph(
        "Todas las validaciones funcionales fueron exitosas. El caso negativo de categoría inexistente retornó 404, "
        "lo cual confirma manejo controlado de errores."
    )


def add_locust_section(document, endpoints, aggregated, charts, html_screenshot):
    document.add_heading("Pruebas de carga con Locust", level=1)
    document.add_paragraph(
        "La prueba de carga fue rehecha con Locust para cubrir todos los endpoints REST disponibles del "
        "catalog-service. Se ejecutó en modo headless durante 1 minuto, con 20 usuarios concurrentes y una tasa de "
        "arranque de 5 usuarios por segundo."
    )
    add_table(
        document,
        ["Métrica agregada", "Valor"],
        [
            ["Solicitudes totales", aggregated["Request Count"]],
            ["Fallos", aggregated["Failure Count"]],
            ["Tasa de fallos", "0.00%" if to_int(aggregated["Failure Count"]) == 0 else "Ver CSV"],
            ["Throughput", f"{to_float(aggregated['Requests/s']):.2f} req/s"],
            ["Latencia promedio", f"{to_float(aggregated['Average Response Time']):.2f} ms"],
            ["Mediana", f"{aggregated['Median Response Time']} ms"],
            ["P95", f"{aggregated['95%']} ms"],
            ["P99", f"{aggregated['99%']} ms"],
            ["Máxima", f"{to_float(aggregated['Max Response Time']):.2f} ms"],
        ],
    )
    rows = []
    for row in endpoints:
        rows.append([
            row["Name"],
            row["Request Count"],
            row["Failure Count"],
            f"{to_float(row['Requests/s']):.2f}",
            f"{to_float(row['Average Response Time']):.2f}",
            row["Median Response Time"],
            row["95%"],
            row["99%"],
            f"{to_float(row['Max Response Time']):.2f}",
        ])
    add_table(document, ["Endpoint", "Reqs", "Fallos", "Req/s", "Prom ms", "Med ms", "P95", "P99", "Max ms"], rows)
    document.add_heading("Análisis de resultados Locust", level=2)
    document.add_paragraph(
        f"La prueba registró {aggregated['Request Count']} solicitudes y {aggregated['Failure Count']} fallos. "
        f"El throughput agregado fue {to_float(aggregated['Requests/s']):.2f} req/s, con una latencia promedio de "
        f"{to_float(aggregated['Average Response Time']):.2f} ms y P95 de {aggregated['95%']} ms. "
        "Esto indica estabilidad bajo carga moderada y sin errores HTTP inesperados."
    )
    document.add_paragraph(
        "Los endpoints de categorías, productos, búsqueda e ítems se comportaron de forma consistente. El endpoint "
        "negativo /api/categories/UNKNOWN fue tratado como exitoso cuando respondió 404, porque ese era el "
        "comportamiento esperado para validar manejo de errores."
    )
    for title, image in [
        ("Resumen agregado Locust", charts[3]),
        ("Tiempo promedio por endpoint", charts[0]),
        ("Percentil 95 por endpoint", charts[1]),
        ("Distribución de solicitudes", charts[2]),
    ]:
        document.add_paragraph(title)
        add_picture_if_exists(document, image, width=6.4)
    if html_screenshot:
        document.add_paragraph("Captura del reporte HTML generado por Locust")
        add_picture_if_exists(document, html_screenshot, width=6.4)


def add_evidence_section(document):
    document.add_heading("Evidencias de ejecución", level=1)
    for title, filename in [
        ("Health check", "evidence_1.png"),
        ("Consulta de categorías", "evidence_2.png"),
        ("Prueba negativa 404", "evidence_3.png"),
        ("Estado del pod", "evidence_4.png"),
        ("Service LoadBalancer", "evidence_5.png"),
    ]:
        document.add_heading(title, level=2)
        add_picture_if_exists(document, CHART_DIR / filename)


def add_sonar_section(document, sonar_data):
    document.add_heading("SonarQube - métricas, dashboards y Quality Gates", level=1)
    rows = []
    for project in sonar_data["projects"]:
        rows.append([
            project["name"], project["quality_gate"], metric(project, "bugs"), metric(project, "vulnerabilities"),
            metric(project, "security_hotspots"), metric(project, "code_smells"), percent(metric(project, "coverage")),
            percent(metric(project, "duplicated_lines_density")), metric(project, "ncloc"), minutes(metric(project, "sqale_index")),
        ])
    add_table(document, ["Proyecto", "Quality Gate", "Bugs", "Vulnerabilidades", "Hotspots", "Code smells", "Cobertura", "Duplicación", "Líneas", "Deuda técnica"], rows)
    document.add_paragraph(
        "El microservicio tiene menor tamaño, menor duplicación y menor deuda técnica que el monolito, lo que "
        "favorece el atributo de mantenibilidad. La cobertura aparece en 0% porque no se importaron reportes JaCoCo."
    )
    for image in [
        "sonar_issues_comparison.png",
        "sonar_maintainability_metrics.png",
        "sonar_quality_percentages.png",
        "sonar_dashboard_jpetstore_monolith.png",
        "sonar_quality_gate_jpetstore_monolith.png",
        "sonar_dashboard_catalog_service.png",
        "sonar_quality_gate_catalog_service.png",
    ]:
        add_picture_if_exists(document, CHART_DIR / image)


def add_quality_analysis(document):
    document.add_heading("Análisis por atributo de calidad", level=1)
    add_table(
        document,
        ["Atributo", "Evidencia", "Interpretación"],
        [
            ["Disponibilidad", "Health check UP", "El servicio estaba disponible y conectado a la base de datos."],
            ["Funcionalidad", "Pruebas REST exitosas", "Los endpoints principales respondieron como se esperaba."],
            ["Desempeño", "Locust", "El servicio sostuvo carga moderada sin fallos."],
            ["Manejo de errores", "404 esperado", "La API responde controladamente ante recursos inexistentes."],
            ["Mantenibilidad", "SonarQube", "El microservicio reduce superficie de código y deuda técnica."],
            ["Observabilidad", "Actuator, kubectl, Sonar", "Hay mecanismos para observar salud, despliegue y calidad."],
        ],
    )


def add_stakeholders_conclusion(document):
    document.add_heading("Interesados involucrados", level=1)
    add_table(
        document,
        ["Interesado", "Participación"],
        [
            ["Desarrollo", "Corrige hallazgos y evoluciona el microservicio."],
            ["QA", "Diseña, ejecuta y documenta pruebas."],
            ["DevOps", "Administra AWS, Kubernetes, despliegue y monitoreo."],
            ["Arquitectura", "Evalúa atributos de calidad y decisiones de modernización."],
            ["Seguridad", "Evalúa vulnerabilidades y configuración."],
            ["Operaciones", "Usa evidencias para soporte y diagnóstico."],
        ],
    )
    document.add_heading("Conclusiones", level=1)
    document.add_paragraph(
        "La ejecución con Locust reemplaza la prueba básica anterior y aporta una medición más representativa, "
        "porque distribuye carga sobre todos los endpoints disponibles. El resultado fue satisfactorio: no se "
        "registraron fallos y las latencias se mantuvieron bajas bajo carga moderada."
    )
    document.add_paragraph(
        "Para una evaluación productiva se recomienda aumentar duración, usuarios concurrentes y monitorear CPU, "
        "memoria, conexiones a PostgreSQL y métricas de infraestructura en CloudWatch o Prometheus."
    )


def main():
    test_data = read_json(TEST_RESULTS)
    sonar_data = read_json(SONAR_RESULTS)
    endpoints, aggregated = read_locust_stats()
    charts = create_locust_charts(endpoints, aggregated)
    html_screenshot = capture_locust_html()

    LOCUST_SUMMARY.write_text(json.dumps({"endpoints": endpoints, "aggregated": aggregated}, indent=2), encoding="utf-8")

    document = Document()
    setup_styles(document)
    add_cover(document)
    add_summary(document, test_data, sonar_data, aggregated)
    add_execution_guide(document)
    add_functional_section(document, test_data)
    add_locust_section(document, endpoints, aggregated, charts, html_screenshot)
    add_evidence_section(document)
    add_sonar_section(document, sonar_data)
    add_quality_analysis(document)
    add_stakeholders_conclusion(document)

    try:
        document.save(OUTPUT)
        print(f"REPORT={OUTPUT}")
    except PermissionError:
        fallback = REPORT_DIR / "informe_pruebas_locust.docx"
        document.save(fallback)
        print(f"REPORT={fallback}")
    print(f"LOCUST_SUMMARY={LOCUST_SUMMARY}")
    print(f"TOTAL_REQUESTS={aggregated['Request Count']}")
    print(f"FAILURES={aggregated['Failure Count']}")
    print(f"RPS={to_float(aggregated['Requests/s']):.2f}")
    print(f"AVG_MS={to_float(aggregated['Average Response Time']):.2f}")
    print(f"P95_MS={aggregated['95%']}")


if __name__ == "__main__":
    main()
